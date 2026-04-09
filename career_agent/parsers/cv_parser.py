"""CV parser — extracts structured data from a PDF or DOCX CV file."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from pydantic import BaseModel


class CVData(BaseModel):
    raw_text: str
    name: str = ""
    email: str = ""
    phone: str = ""
    skills: list[str] = []
    education: str = ""
    projects: list[str] = []
    experience: list[str] = []


def _extract_email(text: str) -> str:
    m = re.search(r"[\w.+-]+@[\w-]+\.[a-zA-Z]{2,}", text)
    return m.group(0) if m else ""


def _extract_phone(text: str) -> str:
    m = re.search(r"(?:\+972|0)[-\s]?\d{2}[-\s]?\d{3}[-\s]?\d{4}", text)
    return m.group(0) if m else ""


def _extract_section(text: str, heading: str) -> str:
    """Return lines between a section heading and the next heading."""
    pattern = rf"(?i){re.escape(heading)}\s*[\n:](.*?)(?=\n[A-Z][A-Z &/]+\s*[\n:]|\Z)"
    m = re.search(pattern, text, re.DOTALL)
    return m.group(1).strip() if m else ""


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ImportError("Install pypdf: pip install pypdf") from e

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _parse_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ImportError("Install python-docx: pip install python-docx") from e

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)


def parse_cv(path: str | Path) -> CVData:
    """Parse a PDF or DOCX CV file and return structured CVData."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"CV file not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw = _parse_pdf(path)
    elif suffix in (".docx", ".doc"):
        raw = _parse_docx(path)
    else:
        raise ValueError(f"Unsupported CV format: {suffix}. Use PDF or DOCX.")

    # Extract first non-empty line as name candidate
    lines = [l.strip() for l in raw.splitlines() if l.strip()]
    name_candidate = lines[0] if lines else ""

    skills_text = _extract_section(raw, "Skills") or _extract_section(raw, "Technical Skills")
    skills = [s.strip() for s in re.split(r"[,\n•\-|]", skills_text) if s.strip()] if skills_text else []

    education = _extract_section(raw, "Education")

    projects_text = _extract_section(raw, "Projects") or _extract_section(raw, "Academic Projects")
    projects = [p.strip() for p in projects_text.split("\n") if p.strip()] if projects_text else []

    experience_text = _extract_section(raw, "Experience") or _extract_section(raw, "Work Experience")
    experience = [e.strip() for e in experience_text.split("\n") if e.strip()] if experience_text else []

    return CVData(
        raw_text=raw,
        name=name_candidate,
        email=_extract_email(raw),
        phone=_extract_phone(raw),
        skills=skills[:30],    # cap to avoid noise
        education=education,
        projects=projects[:10],
        experience=experience[:20],
    )
