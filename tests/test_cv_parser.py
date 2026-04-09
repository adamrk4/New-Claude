"""Tests for the CV parser module."""
from __future__ import annotations

import tempfile
from pathlib import Path

import pytest

from career_agent.parsers.cv_parser import _extract_email, _extract_phone, parse_cv


class TestExtractEmail:
    def test_simple_email(self):
        assert _extract_email("Contact: john@example.com") == "john@example.com"

    def test_no_email(self):
        assert _extract_email("No email here") == ""

    def test_multiple_emails_first_wins(self):
        result = _extract_email("a@b.com and c@d.com")
        assert result == "a@b.com"


class TestExtractPhone:
    def test_israeli_mobile(self):
        result = _extract_phone("Phone: 050-1234567")
        assert "050" in result

    def test_international(self):
        result = _extract_phone("Tel: +972-50-1234567")
        assert "+972" in result

    def test_no_phone(self):
        assert _extract_phone("No phone here") == ""


class TestParseCv:
    def test_unsupported_format_raises(self):
        with tempfile.NamedTemporaryFile(suffix=".txt") as f:
            with pytest.raises(ValueError, match="Unsupported CV format"):
                parse_cv(f.name)

    def test_missing_file_raises(self):
        with pytest.raises(FileNotFoundError):
            parse_cv("/nonexistent/cv.pdf")

    def test_parse_docx(self, tmp_path):
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "cv.docx"
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("jane@example.com")
        doc.add_paragraph("050-9876543")
        doc.add_paragraph("Education\nB.Sc. Computer Science, Tel Aviv University, 2024")
        doc.add_paragraph("Skills\nPython, SQL, REST APIs, Communication")
        doc.save(str(doc_path))

        cv = parse_cv(doc_path)
        assert cv.email == "jane@example.com"
        assert "050" in cv.phone
        assert cv.raw_text  # not empty
